from openpyxl import load_workbook
from pathlib import Path
from modules.categorizer import apply_categories
from modules.import_review import review_transactions
from modules.duplicate_detector import (
    filter_duplicates,
    print_duplicate_summary
)
from modules.statement_import import save_imported_transactions
from modules.duplicate_detector import register_imported_transactions
def safe_float(value):
    """
    Convert an Excel cell value into float.

    Returns None when the cell is empty or invalid.
    """
    if value is None or value == "":
        return None

    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def find_transaction_header(sheet):
    """
    Find the row containing the IPPB transaction table header.

    Expected columns:
    DATE
    TRAN ID
    TRANSACTION PARTICULARS
    WITHDRWAL
    DEPOSIT
    BALANCE
    """

    for row_number in range(1, sheet.max_row + 1):

        date_value = sheet.cell(row_number, 1).value
        transaction_id_value = sheet.cell(row_number, 2).value
        particulars_value = sheet.cell(row_number, 3).value

        if (
            str(date_value).strip().upper() == "DATE"
            and str(transaction_id_value).strip().upper() == "TRAN ID"
            and "TRANSACTION PARTICULARS"
            in str(particulars_value).strip().upper()
        ):
            return row_number

    return None


def parse_ippb_excel(file_path):
    """
    Read an IPPB Excel bank statement and return
    normalized InsightSpend transactions.

    Returns:
        list[dict]
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Bank statement not found: {file_path}"
        )

    if file_path.suffix.lower() not in [".xlsx", ".xlsm"]:
        raise ValueError(
            "IPPB parser currently supports .xlsx or .xlsm files only."
        )

    workbook = load_workbook(
        file_path,
        data_only=True
    )

    sheet = workbook.active

    header_row = find_transaction_header(sheet)

    if header_row is None:
        raise ValueError(
            "Could not find the IPPB transaction table."
        )

    transactions = []

    # Transaction data starts after:
    # header row
    # opening balance row
    start_row = header_row + 2

    for row_number in range(start_row, sheet.max_row + 1):

        date = sheet.cell(row_number, 1).value
        transaction_id = sheet.cell(row_number, 2).value
        description = sheet.cell(row_number, 3).value
        withdrawal = sheet.cell(row_number, 4).value
        deposit = sheet.cell(row_number, 5).value
        balance = sheet.cell(row_number, 6).value

        # Stop when account summary begins
        if str(date).strip().upper() == "ACCOUNT SUMMARY":
            break

        # Ignore empty rows
        if not date and not transaction_id:
            continue

        withdrawal = safe_float(withdrawal)
        deposit = safe_float(deposit)
        balance = safe_float(balance)

        # Determine transaction type
        if withdrawal is not None and withdrawal > 0:
            transaction_type = "expense"
            amount = withdrawal

        elif deposit is not None and deposit > 0:
            transaction_type = "income"
            amount = deposit

        else:
            # Ignore rows without an actual transaction amount
            continue

        transaction = {
            "date": str(date).strip(),
            "transaction_id": str(transaction_id).strip(),
            "description": str(description).strip(),
            "amount": amount,
            "type": transaction_type,
            "category": "Uncategorized",
            "balance": balance,
            "source": "IPPB",
            "source_type": "bank_statement"
        }

        transactions.append(transaction)

    return transactions


def print_transactions(transactions):
    """
    Display imported transactions for testing.
    """

    print("\n" + "=" * 70)
    print("              IPPB BANK STATEMENT IMPORT")
    print("=" * 70)

    if not transactions:
        print("\nNo transactions found.")
        return

    total_income = 0
    total_expense = 0

    for index, transaction in enumerate(transactions, start=1):

        amount = transaction["amount"]
        transaction_type = transaction["type"]

        if transaction_type == "income":
            total_income += amount
            symbol = "+"

        else:
            total_expense += amount
            symbol = "-"

        print(f"\nTransaction #{index}")
        print("-" * 40)

        print(f"Date        : {transaction['date']}")
        print(f"Transaction : {transaction['transaction_id']}")
        print(f"Description : {transaction['description']}")
        print(f"Amount      : {symbol}₹{amount:.2f}")
        print(f"Type        : {transaction_type.title()}")
        print(f"Category    : {transaction['category']}")
        if "category_confidence" in transaction:
            print(
            f"Confidence  : {transaction['category_confidence']}"
            )

        if "category_reason" in transaction:
            print(
                f"Reason      : {transaction['category_reason']}"
                )
        if transaction["balance"] is not None:
            print(
                f"Balance     : ₹{transaction['balance']:.2f}"
            )

    print("\n" + "=" * 70)
    print("SUMMARY")
    print("=" * 70)

    print(f"Transactions   : {len(transactions)}")
    print(f"Total Income   : ₹{total_income:.2f}")
    print(f"Total Expenses : ₹{total_expense:.2f}")
    print(
        f"Net Change     : ₹{total_income - total_expense:.2f}"
    )

    print("=" * 70)

if __name__ == "__main__":

    statement_path = input(
        "Enter path of IPPB Excel statement: "
    ).strip().strip('"')

    try:
        transactions = parse_ippb_excel(
            statement_path
        )

        transactions = apply_categories(
            transactions
        )

        approved_transactions = review_transactions(
            transactions
        )

        new_transactions, duplicate_transactions = (
            filter_duplicates(
                approved_transactions
            )
        )

        print_duplicate_summary(
            new_transactions,
            duplicate_transactions
        )

        print("\nApproved transactions ready for import:")

        for transaction in new_transactions:
            print(
                transaction["date"],
                transaction["type"],
                transaction["category"],
                transaction["amount"]
            )
       
        saved_transactions = save_imported_transactions(
            new_transactions
        )

        if saved_transactions:
            registered_count = register_imported_transactions(
                saved_transactions
            )

            print(
                f"\nRegistered {registered_count} "
                f"transactions in import history."
            )
        else:
            print("\nNo transactions were saved.")

    except FileNotFoundError as error:
        print(f"\nError: {error}")

    except ValueError as error:
        print(f"\nInvalid statement: {error}")

    except Exception as error:
        print(
            f"\nUnexpected error while reading statement: {error}"
        )
